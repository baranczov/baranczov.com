from pathlib import Path
import json, re
from collections import Counter
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parent
rows = json.loads((ROOT/'selected-draft.json').read_text())
rows = [r for r in rows if r['id'] not in {27,70,135}]
notes = {
 1:'LLM agents; от 3 лет опыта.', 2:'AI для private markets; работа в офисе 4 дня в неделю.',
 4:'Описание на испанском; аналитика маркетинга.', 6:'В описании указан срок 6 месяцев; уточнить формат найма.',
 7:'NLP и conversational AI; от 2 лет.', 8:'Указана помощь с переездом; визу уточнить отдельно.',
 9:'LLM applications и работа с корпоративными клиентами.', 10:'ML для legal tech.',
 11:'От 3 лет ML и MLOps; удалённо внутри США.',12:'AI в fintech; гибрид.',
 13:'От 3 лет; SQL, Python, DWH и ETL.',14:'От 1,5 лет ML/DL; n8n, vLLM, Qdrant.',
 15:'От 4 лет Python; FastAPI и LangChain.',16:'AI разработка; в карточке Junior, в названии Junior/Senior.',
 17:'RAG и AI agents; Python от 2 лет.',18:'Big Data, ETL, Python, SQL и Greenplum.',
 19:'Локация и гражданство РФ; Python, SQL и LLM.',21:'Прикладные AI/ML, LLM, RAG и CV.',
 22:'Python от 1 года; ML и оптимизация в e-commerce.',25:'Python, SQL и статистические модели.',
 29:'Начальная позиция; офис.',30:'От 3 лет Data Science или Advanced Analytics.',
 31:'Associate; начальная позиция в Data Analytics and AI.',34:'Удалённый формат; география указана как Берлин.',
 35:'Аналитика производственных процессов.',38:'Английский; SQL и бизнес аналитика.',
 39:'Аналитика продукта и бизнеса.',40:'Аналитика в продуктовой компании; офис.',41:'Аналитика платформы занятости.',
 43:'В форме есть вопрос о спонсорстве; это не гарантия визы.',45:'1–3 года аналитики; удалённо внутри США.',
 46:'Python и SQL; банковские данные.',47:'Количественные исследования и ML; агентство.',
 48:'Рекомендательные системы; в описании Senior.',49:'ML в металлургической отрасли.',50:'От 2 лет анализа данных и ML.',
 51:'От 2 лет DS/ML.',52:'Python, SQL и прикладное моделирование; агентство.',
 53:'Частичная занятость; от 3–4 лет.',55:'DS в разработке лекарств; возможны командировки.',
 57:'От 6 лет; персонализация, ranking и causal measurement.',58:'От 5 лет; маркетинговые модели.',
 59:'Требуются магистратура и 3–4 года либо PhD и 2+ года.',63:'От 5 лет; marketing measurement.',
 64:'От 4 лет; remote только США или Канада.',65:'От 5 лет; DNN, bidding и прогнозирование.',
 66:'От 8 лет с bachelor либо от 10 лет опыта.',71:'3–5 лет ML; AWS и production.',
 72:'Physics AI; вакансию размещает рекрутер.',73:'VLA, PyTorch/JAX; степень или эквивалентный опыт.',
 74:'Название Senior/Lead; рассматривать уровень Senior.',75:'Оценка качества ML моделей.',
 78:'Anti-fraud ML; список допустимых стран в карточке.',79:'LLM и ML разработка.',81:'Разработка новых AI продуктов.',
 82:'Обязателен японский на уровне носителя; визы спонсируют с оговорками.',
 83:'Applied AI и внедрение решений у клиентов.',86:'Applied AI для coding tools и Codex.',
 87:'Senior AI; гибрид.',88:'Agent runtime и AI приложения.',90:'Applied AI в cybersecurity.',
 92:'Список допустимых стран в карточке; РФ не указана.',95:'В форме есть вопрос о визе; поддержку уточнить.',
 97:'2–3+ года DA/BA; remote в перечисленных странах Азии.',99:'В тексте Junior; при этом требуется от 2 лет.',
 100:'Junior AI — название компании; грейд вакансии не указан.',102:'Energy analytics; описание на немецком.',
 103:'1–2 года DA/DS; нужен английский.',104:'От 5 лет; в форме вопрос о визе.',105:'AI в медицине; удалённо во Франции.',
 107:'Python, PostgreSQL и ML.',108:'В описании Middle+/Senior; Python от 3 лет.',109:'GenAI tooling для разработки.',
 113:'От 3 лет DS/ML и от 2 лет рекомендательных систем.',116:'AI исследования; агентство.',117:'ML, PyTorch, LangGraph и CI/CD.',
 118:'SQL и анализ данных.',119:'Белград и remote; уточнить допустимую географию.',120:'SQL и анализ данных cross-border.',
 121:'BI и аналитика пунктов выдачи.',122:'BI, SQL и Python.',123:'Data Science в банке.',124:'SQL и банковская аналитика.',
 125:'Junior; офис в Нью-Йорке.',126:'От 1 года; прямо указаны рабочая виза и помощь с переездом.',
 127:'В описании диапазон Junior–Mid–Senior.',131:'1–3 года; требуется резюме на английском.',
 132:'От 3 лет; объявление от посредника, работодатель не раскрыт.',137:'Обычно 1–2 года; допускается сильный опыт стажировок.',
 138:'Junior DS для банковской команды; визовый вопрос в форме.',139:'Требуется оконченная магистратура; проживание возле Барселоны.',
 140:'Прямо указаны виза и релокация; требуется bachelor/master.',142:'Applied AI; описание на французском.',
 143:'GenAI и LLM в production; описание на французском.',144:'AI разработка; описание на французском.',
 145:'Senior AI; гибрид.',146:'Антифрод аналитика в играх; английский B2+.',148:'От 3 лет продуктовой аналитики; iGaming.',
 150:'От 4 лет Python; AWS, ML и NLP.'
}

def clean_page(r):
    return re.sub(r'^L\d+: ?', '', r['raw_page'], flags=re.M)

for r in rows:
    r['note'] = notes.get(r['id'],'Условия и требования в карточке.')
    r['location'] = r['location'].replace('United States','США').replace('Remote','удалённо').replace('remote','удалённо')
    if 'career.habr.com' in r['url']:
        m=re.search(r'## Условия\n(.*?)\n## Компания',clean_page(r),re.S)
        lines=[s.strip() for s in m.group(1).splitlines() if s.strip() and 'Отклик' not in s] if m else []
        remote='Можно удалённо' in lines
        cities=[s for s in lines if s not in ['Можно удалённо','Неполный рабочий день']]
        country='Сербия' if r['id']==119 else 'Россия'
        r['location']=country+(' · '+', '.join(cities) if cities else '')+(' · удалённо' if remote else '')
    else:
        m=re.search(r'## Location Type\s*\n\s*([^\n]+)',r['raw_search'])
        if m:
            fmt={'Hybrid':'гибрид','On-site':'офис','Remote':'удалённо'}.get(m.group(1).strip(),'')
            if fmt and fmt not in r['location']:r['location']+=' · '+fmt
    if r['id'] in {21,22}:r['location']='Россия · формат уточнить'
    if r['id']==64:r['location']='США или Канада · удалённо'
    if r['id']==74:r['location']='Европа и Канада · удалённо; точные страны по ссылке'
    if r['id']==78:r['location']='Великобритания и ряд стран Европы · удалённо'
    if r['id']==92:r['location']='Латинская Америка, Африка, Азия · удалённо; список по ссылке'
    if r['id']==95:r['location']='Сингапур · гибрид'
    if r['id']==97:r['location']='Индия, Сингапур, Таиланд, Филиппины, Вьетнам · удалённо'
    if r['id']==105:r['location']='Франция · несколько городов или удалённо'
    if r['id']==113:r['location']='Россия · удалённо';r['company']='2ГИС'
    if r['id']==145:r['location']='Франция или Греция · Париж или Афины · гибрид'
    if r['id']==99:r['grade']='Junior в описании'
    if r['id']==127:r['grade']='Junior–Senior'
    if r['id']==108:r['grade']='Middle–Senior'
    if r['id']==131:r['grade']='Не указан · 1–3 года'
    if r['id']==22:r['grade']='Не указан · от 1 года'
    if r['id']==7:r['location']='Испания · удалённо'
    if r['company']=='MAGNIT TECH':r['company']='Магнит Тех'
    if r['company']=='Cobblestone Energy.':r['company']='Cobblestone Energy'
    if r['id'] in {82,1,2,6,8,9,12,14,15,16,17,21,29,81,83,86,87,88,90,108,109,116,142,143,144,145}:r['category']='AIE'
    elif r['id'] in {7,10,11,22,48,71,72,73,74,75,78,79,107,117,125,126,127}:r['category']='MLE'
    elif re.search(r'analyst|аналитик',r['role'],re.I):r['category']='DA'
    else:r['category']='DS'

assert len(rows)==100
assert len({r['url'] for r in rows})==100
rank=lambda r: (0 if 'Junior' in r['grade'] or 'Associate' in r['role'] else 1 if 'Middle' in r['grade'] else 3 if r['grade']=='Senior' else 2, 0 if r['location'].startswith('Россия') else 1, r['company'].casefold())
rows.sort(key=lambda r: (['MLE','AIE','DS','DA'].index(r['category']),rank(r)))
for n,r in enumerate(rows,1):r['number']=n
stats=Counter(r['category'] for r in rows)
checks=Counter(r['status'] for r in rows)
(ROOT/'vacancies.json').write_text(json.dumps(rows,ensure_ascii=False,indent=2)+'\n')

doc=Document()
sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.5);sec.bottom_margin=Cm(1.5);sec.left_margin=Cm(1.35);sec.right_margin=Cm(1.35)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
    st=doc.styles[name];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
    st.element.get_or_add_rPr().append(OxmlElement('w:rFonts')) if st.element.get_or_add_rPr().find(qn('w:rFonts')) is None else None
    rf=st.element.get_or_add_rPr().find(qn('w:rFonts'));rf.set(qn('w:ascii'),'Arial');rf.set(qn('w:hAnsi'),'Arial')
doc.styles['Normal'].font.size=Pt(10)
doc.styles['Normal'].paragraph_format.space_after=Pt(5)
doc.styles['Normal'].paragraph_format.line_spacing=1.08
doc.styles['Title'].font.size=Pt(25)
doc.styles['Title'].paragraph_format.space_after=Pt(10)
doc.styles['Heading 1'].font.size=Pt(17)
doc.styles['Heading 1'].paragraph_format.space_before=Pt(14)
doc.styles['Heading 1'].paragraph_format.space_after=Pt(8)
doc.styles['Heading 2'].font.size=Pt(12)
doc.core_properties.title='Вакансии MLE AIE DS DA для самостоятельных откликов'
doc.core_properties.subject='Прямые ссылки на вакансии от Junior до Senior'
doc.core_properties.author='Данил Баранцов'

def link(p,label,url):
    h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
    run=OxmlElement('w:r');pr=OxmlElement('w:rPr');col=OxmlElement('w:color');col.set(qn('w:val'),'175CD3');pr.append(col)
    u=OxmlElement('w:u');u.set(qn('w:val'),'single');pr.append(u)
    sz=OxmlElement('w:sz');sz.set(qn('w:val'),'19');pr.append(sz);run.append(pr)
    text=OxmlElement('w:t');text.text=label;run.append(text);h.append(run);p._p.append(h)

def table(headers,widths):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for c,w,h in zip(t.rows[0].cells,widths,headers):
        c.width=Cm(w);c.text=h
        for run in c.paragraphs[0].runs:run.bold=True;run.font.color.rgb=RGBColor(255,255,255);run.font.size=Pt(9)
        sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'263D55');c._tc.get_or_add_tcPr().append(sh)
    trpr=t.rows[0]._tr.get_or_add_trPr();repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
    borders=OxmlElement('w:tblBorders')
    for name in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+name);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    t._tbl.tblPr.append(borders)
    for col,w in zip(t.columns,widths):col.width=Cm(w)
    return t

def format_row(row, index):
    pr=row._tr.get_or_add_trPr();cs=OxmlElement('w:cantSplit');pr.append(cs)
    for cell in row.cells:
        cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        mar=OxmlElement('w:tcMar')
        for side,n in [('top',85),('bottom',85),('left',90),('right',90)]:
            e=OxmlElement('w:'+side);e.set(qn('w:w'),str(n));e.set(qn('w:type'),'dxa');mar.append(e)
        cell._tc.get_or_add_tcPr().append(mar)
        if index%2==0:
            sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F2F5F8');cell._tc.get_or_add_tcPr().append(sh)
        for p in cell.paragraphs:
            p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1.03
            for run in p.runs:run.font.size=Pt(9.5)

doc.add_paragraph('Вакансии MLE AIE DS DA',style='Title')
doc.add_paragraph('100 прямых ссылок для самостоятельных откликов',style='Subtitle')
doc.add_paragraph('Дата сборки и проверки 26 октября 2026 года')
doc.add_paragraph(f'В подборке {stats["MLE"]} вакансий Machine Learning Engineer, {stats["AIE"]} вакансий AI Engineer, {stats["DS"]} вакансий Data Scientist и {stats["DA"]} вакансий Data Analyst. Есть позиции Junior, Middle и Senior, а также вакансии без явно указанного грейда. Источники — карьерные сайты, страницы работодателей в ATS и Хабр Карьера.')
doc.add_paragraph('Для отклика нажмите «Открыть вакансию» в нужной строке. В последнем столбце можно записать дату отправки и ответ работодателя. Ссылки ведут на конкретные объявления.')
doc.add_paragraph('Как читать проверку',style='Heading 2')
doc.add_paragraph(f'Карточка — прочитана страница конкретной вакансии ({checks["Карточка"]} ссылок). Индекс — описание найдено в поисковом индексе, но актуальный текст динамической страницы получить не удалось ({checks["Индекс"]} ссылок); проверьте, что форма ещё принимает отклики. Пометки относятся к дате сборки, вакансии могут закрыться позже.')
doc.add_paragraph('Как читать условия',style='Heading 2')
doc.add_paragraph('«Не указан» означает, что работодатель не назвал грейд прямо. Указанный стаж приведён отдельно, где он найден. Удалённая работа часто ограничена перечисленными странами. Визовая поддержка подтверждена только там, где это прямо отмечено; в остальных случаях её нужно уточнить у работодателя.')
doc.add_paragraph('Для приоритизации',style='Heading 2')
doc.add_paragraph('Внутри каждого направления сначала идут Junior и Middle. У Revel Robotics прямо указаны рабочая виза и помощь с переездом. У Cobblestone Energy также указаны виза и релокация, но требуется завершённая степень. Для Buynomics требуется оконченная магистратура. У Anthropic в Токио обязателен японский на уровне носителя.')
doc.add_paragraph('Ранее отправленные заявки и неподтверждённая попытка Lamoda вынесены в конец документа. Они не входят в 100 вакансий основного списка.')

headings={'MLE':'MLE Machine Learning Engineer','AIE':'AIE AI Engineer','DS':'DS Data Scientist','DA':'DA Data Analyst'}
for cat in ['MLE','AIE','DS','DA']:
    if cat=='MLE':doc.add_page_break()
    doc.add_paragraph(headings[cat],style='Heading 1')
    subset=[r for r in rows if r['category']==cat]
    intro=doc.add_paragraph(f'{len(subset)} вакансий. Строки с пометкой «Индекс» требуют дополнительной проверки перед заполнением формы.')
    intro.paragraph_format.keep_with_next=True
    t=table(['№','Компания и вакансия','Грейд','География и формат','Условия и мой отклик'],[1.0,5.65,2.1,3.75,5.8])
    for k,r in enumerate(subset):
        cells=t.add_row().cells
        cells[0].text=str(r['number']);cells[0].paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER
        cells[0]._tc.get_or_add_tcPr().append(OxmlElement('w:noWrap'))
        cells[1].paragraphs[0].add_run(r['company']).bold=True
        cells[1].add_paragraph(r['role'])
        link(cells[1].add_paragraph(),'Открыть вакансию',r['url'])
        cells[2].text=r['grade'];cells[3].text=r['location']
        cells[4].paragraphs[0].add_run(r['status']+' · 26.10.2026').bold=True
        cells[4].add_paragraph(r['note'])
        p=cells[4].add_paragraph('Отклик ______  Ответ ______');p.paragraph_format.space_before=Pt(3)
        format_row(t.rows[-1],k)

doc.add_page_break()
doc.add_paragraph('Ранее отправленные отклики',style='Heading 1')
doc.add_paragraph('Четыре отправки подтверждены страницами успеха. По Lamoda была нажата кнопка отправки, но подтверждение прочитать не удалось. Перед повторной подачей проверьте почту и историю откликов.')
apps=json.loads((ROOT.parent/'applications/2026-10-26/applications.json').read_text())['applications']
prior=[a for a in apps if a['status'] in ['submitted_confirmed','submission_attempted_unconfirmed']]
t=table(['Компания и вакансия','Результат'],[11.0,7.3])
for k,a in enumerate(prior):
    c=t.add_row().cells;c[0].paragraphs[0].add_run(a['company']).bold=True;c[0].add_paragraph(a['role']);link(c[0].add_paragraph(),'Открыть вакансию',a['url'])
    c[1].text='Отправлено и подтверждено' if a['status']=='submitted_confirmed' else 'Попытка отправки · результат не подтверждён'
    format_row(t.rows[-1],k)

footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
run=footer.add_run('Вакансии · 26 октября 2026 · ');run.font.size=Pt(8)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
out=ROOT/'Вакансии_MLE_AIE_DS_DA_2026-10-26.docx'
for tree in [doc._element,doc.styles.element]:
    for border in list(tree.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
doc.save(out)
print(json.dumps({'path':str(out),'count':len(rows),'categories':dict(stats),'verification':dict(checks),'prior_applications':len(prior)},ensure_ascii=False))
